import tempfile
import unittest
from pathlib import Path

from docx import Document

from app.parser import parse_docx, parse_document, parse_pdf


def _make_pdf(page_texts: list[str]) -> bytes:
    """构造一个最小、可被 pypdf 读取的 PDF，每页含一段文本。"""
    count = len(page_texts)
    catalog_id, pages_id = 1, 2
    page_ids = list(range(3, 3 + count))
    content_ids = list(range(3 + count, 3 + 2 * count))
    font_id = 3 + 2 * count

    objects: dict[int, bytes] = {}
    objects[catalog_id] = f"<< /Type /Catalog /Pages {pages_id} 0 R >>".encode("ascii")
    kids = " ".join(f"{i} 0 R" for i in page_ids)
    objects[pages_id] = f"<< /Type /Pages /Kids [{kids}] /Count {count} >>".encode("ascii")

    for page_id, content_id in zip(page_ids, content_ids):
        objects[page_id] = (
            f"<< /Type /Page /Parent {pages_id} 0 R /MediaBox [0 0 612 792] "
            f"/Contents {content_id} 0 R /Resources << /Font << /F1 {font_id} 0 R >> >> >>"
        ).encode("ascii")

    for content_id, text in zip(content_ids, page_texts):
        escaped = text.replace("\\", "\\\\").replace("(", "\\(").replace(")", "\\)")
        stream = f"BT /F1 12 Tf 72 720 Td ({escaped}) Tj ET".encode("ascii")
        objects[content_id] = (
            f"<< /Length {len(stream)} >>\nstream\n".encode("ascii")
            + stream
            + b"\nendstream"
        )

    objects[font_id] = b"<< /Type /Font /Subtype /Type1 /BaseFont /Helvetica >>"

    output = bytearray(b"%PDF-1.4\n")
    offsets: dict[int, int] = {}
    for object_id in range(1, font_id + 1):
        offsets[object_id] = len(output)
        output += f"{object_id} 0 obj\n".encode("ascii")
        output += objects[object_id]
        output += b"\nendobj\n"

    xref_position = len(output)
    output += f"xref\n0 {font_id + 1}\n".encode("ascii")
    output += b"0000000000 65535 f \n"
    for object_id in range(1, font_id + 1):
        output += f"{offsets[object_id]:010d} 00000 n \n".encode("ascii")
    output += (
        f"trailer\n<< /Size {font_id + 1} /Root {catalog_id} 0 R >>\n"
        f"startxref\n{xref_position}\n%%EOF\n"
    ).encode("ascii")
    return bytes(output)


class ParserTests(unittest.TestCase):
    def test_parse_pdf_extracts_text_per_page(self):
        pdf = _make_pdf(["Hello Page One", "Hello Page Two"])
        with tempfile.TemporaryDirectory() as tmp:
            path = Path(tmp) / "cv.pdf"
            path.write_bytes(pdf)
            parsed = parse_pdf(path)

        self.assertEqual([page.page_number for page in parsed.pages], [1, 2])
        self.assertIn("Hello Page One", parsed.pages[0].text)
        self.assertIn("Hello Page Two", parsed.pages[1].text)
        self.assertEqual(parsed.content_type, "application/pdf")

    def test_parse_docx_extracts_paragraphs_and_tables(self):
        with tempfile.TemporaryDirectory() as tmp:
            path = Path(tmp) / "cv.docx"
            doc = Document()
            doc.add_paragraph("姓名：张三")
            table = doc.add_table(rows=1, cols=2)
            table.rows[0].cells[0].text = "博士"
            table.rows[0].cells[1].text = "计算机科学"
            doc.save(path)
            parsed = parse_docx(path)

        self.assertIn("姓名：张三", parsed.full_text())
        self.assertIn("博士", parsed.full_text())
        self.assertIn("计算机科学", parsed.full_text())
        self.assertIsNone(parsed.pages[0].page_number)

    def test_parse_document_dispatches_by_extension(self):
        with tempfile.TemporaryDirectory() as tmp:
            path = Path(tmp) / "cv.pdf"
            path.write_bytes(_make_pdf(["Page"]))
            parsed = parse_document(path)
        self.assertEqual(parsed.content_type, "application/pdf")


if __name__ == "__main__":
    unittest.main()
