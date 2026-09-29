"""文档解析（D2）：PDF / DOCX -> 纯文本 + 页码。

页码用于证据定位（schema.json 中 evidence.page）。PDF 逐页提取文本并保留真实
页码；DOCX 无法可靠获得分页信息，因此整篇返回为单条、页码为 None，上层生成
证据时不填 page 字段。
"""

from __future__ import annotations

from dataclasses import dataclass
from pathlib import Path

from docx import Document
from pypdf import PdfReader


SUPPORTED_SUFFIXES = {".pdf", ".docx"}

CONTENT_TYPES = {
    ".pdf": "application/pdf",
    ".docx": "application/vnd.openxmlformats-officedocument.wordprocessingml.document",
}


@dataclass
class PageText:
    """一页（或一个逻辑文本段）的提取结果。"""

    page_number: int | None
    text: str


@dataclass
class ParsedDocument:
    """解析后的文档。"""

    filename: str
    content_type: str
    pages: list[PageText]

    def full_text(self) -> str:
        return "\n".join(page.text for page in self.pages if page.text)


def _normalize_text(text: str | None) -> str:
    return (text or "").strip()


def parse_pdf(path: str | Path) -> ParsedDocument:
    path = Path(path)
    reader = PdfReader(str(path))
    pages = [
        PageText(page_number=index, text=_normalize_text(page.extract_text()))
        for index, page in enumerate(reader.pages, start=1)
    ]
    return ParsedDocument(
        filename=path.name,
        content_type=CONTENT_TYPES[".pdf"],
        pages=pages,
    )


def parse_docx(path: str | Path) -> ParsedDocument:
    path = Path(path)
    doc = Document(str(path))

    chunks: list[str] = []
    for paragraph in doc.paragraphs:
        if paragraph.text.strip():
            chunks.append(paragraph.text.strip())
    # 简历常使用表格排版，遍历表格单元格并合并为一行文本。
    for table in doc.tables:
        for row in table.rows:
            cells = [cell.text.strip() for cell in row.cells if cell.text.strip()]
            if cells:
                chunks.append(" | ".join(cells))

    return ParsedDocument(
        filename=path.name,
        content_type=CONTENT_TYPES[".docx"],
        pages=[PageText(page_number=None, text="\n".join(chunks))],
    )


def parse_document(path: str | Path) -> ParsedDocument:
    """按扩展名分派到 PDF 或 DOCX 解析器。"""
    path = Path(path)
    suffix = path.suffix.lower()
    if suffix == ".pdf":
        return parse_pdf(path)
    if suffix == ".docx":
        return parse_docx(path)
    raise ValueError(f"不支持的文档类型：{suffix or '（无扩展名）'}")
