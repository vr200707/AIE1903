from __future__ import annotations

import argparse
import csv
import hashlib
import json
from pathlib import Path
from typing import Any

from docx import Document
from pypdf import PdfReader


PROJECT_ROOT = Path(__file__).resolve().parents[1]
DEFAULT_SAMPLE_DIR = (
    PROJECT_ROOT / "data" / "samples" / "cv-fixtures-2026-09-29"
)


def sha256(path: Path) -> str:
    return hashlib.sha256(path.read_bytes()).hexdigest()


def inspect_pdf(path: Path) -> dict[str, Any]:
    reader = PdfReader(str(path))
    pages = [page.extract_text() or "" for page in reader.pages]
    return {
        "page_count": len(reader.pages),
        "text_characters": sum(len(text) for text in pages),
        "empty_text_pages": sum(not text.strip() for text in pages),
    }


def inspect_docx(path: Path) -> dict[str, Any]:
    document = Document(str(path))
    parts = [paragraph.text for paragraph in document.paragraphs]
    for table in document.tables:
        for row in table.rows:
            parts.extend(cell.text for cell in row.cells)
    text = "\n".join(part for part in parts if part)
    return {
        "paragraph_count": len(document.paragraphs),
        "table_count": len(document.tables),
        "text_characters": len(text),
    }


def inspect_file(sample_dir: Path, row: dict[str, str]) -> dict[str, Any]:
    path = sample_dir / row["filename"]
    issues: list[str] = []
    result: dict[str, Any] = {
        "filename": row["filename"],
        "format": row["format"],
        "expected_text_layer": row["text_layer"],
        "issues": issues,
    }

    if not path.is_file():
        issues.append("file is missing")
        return result

    actual_size = path.stat().st_size
    actual_hash = sha256(path)
    result["size_bytes"] = actual_size
    result["sha256"] = actual_hash
    if actual_size != int(row["size_bytes"]):
        issues.append("size does not match manifest")
    if actual_hash.lower() != row["sha256"].lower():
        issues.append("SHA-256 does not match manifest")

    suffix = path.suffix.lower()
    try:
        if suffix == ".pdf":
            details = inspect_pdf(path)
        elif suffix == ".docx":
            details = inspect_docx(path)
        else:
            issues.append(f"unsupported extension: {suffix}")
            return result
    except Exception as error:
        issues.append(f"cannot read: {error}")
        return result

    result.update(details)
    if suffix == ".pdf" and details["page_count"] != int(row["pages"]):
        issues.append("page count does not match manifest")

    has_text = details["text_characters"] > 0
    expects_text = row["text_layer"].lower() == "yes"
    if has_text != expects_text:
        issues.append("text-layer expectation does not match")

    return result


def main() -> int:
    parser = argparse.ArgumentParser(
        description="Check synthetic PDF/DOCX sample readability and integrity."
    )
    parser.add_argument(
        "--samples",
        type=Path,
        default=DEFAULT_SAMPLE_DIR,
        help="Directory containing manifest.csv and sample documents.",
    )
    parser.add_argument(
        "--json",
        action="store_true",
        help="Print machine-readable JSON instead of a text table.",
    )
    args = parser.parse_args()

    manifest_path = args.samples / "manifest.csv"
    if not manifest_path.is_file():
        print(f"manifest not found: {manifest_path}")
        return 2

    with manifest_path.open(encoding="utf-8-sig", newline="") as handle:
        rows = list(csv.DictReader(handle))

    results = [inspect_file(args.samples, row) for row in rows]
    issue_count = sum(len(result["issues"]) for result in results)

    if args.json:
        print(
            json.dumps(
                {"sample_dir": str(args.samples), "results": results},
                ensure_ascii=False,
                indent=2,
            )
        )
    else:
        for result in results:
            status = "OK" if not result["issues"] else "FAIL"
            text_chars = result.get("text_characters", "-")
            pages = result.get("page_count", "-")
            print(
                f"{status:4} {result['filename']:38} "
                f"pages={pages:<3} text_chars={text_chars}"
            )
            for issue in result["issues"]:
                print(f"     - {issue}")
        print(f"SUMMARY files={len(results)} issues={issue_count}")

    return 1 if issue_count else 0


if __name__ == "__main__":
    raise SystemExit(main())
