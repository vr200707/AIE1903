#!/usr/bin/env python3
"""Generate synthetic CV files for parser and extraction testing."""

from __future__ import annotations

import csv
import hashlib
import os
import random
import shutil
from pathlib import Path
from typing import Iterable, Sequence

from docx import Document
from docx.enum.section import WD_SECTION
from docx.enum.table import WD_CELL_VERTICAL_ALIGNMENT, WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Cm, Inches, Pt, RGBColor
from PIL import Image, ImageEnhance, ImageOps
from pdf2image import convert_from_path
from pypdf import PdfReader
from reportlab.lib import colors
from reportlab.lib.enums import TA_CENTER, TA_LEFT, TA_RIGHT
from reportlab.lib.pagesizes import A4, letter
from reportlab.lib.styles import ParagraphStyle
from reportlab.lib.units import mm
from reportlab.pdfbase import pdfmetrics
from reportlab.pdfbase.ttfonts import TTFont
from reportlab.pdfgen import canvas
from reportlab.platypus import (
    BaseDocTemplate,
    Frame,
    HRFlowable,
    KeepTogether,
    PageBreak,
    PageTemplate,
    Paragraph,
    SimpleDocTemplate,
    Spacer,
    Table,
    TableStyle,
)


ROOT = Path(__file__).resolve().parents[2]
OUT_DIR = ROOT / "data" / "samples" / "cv-fixtures-2026-09-29"
DEFAULT_FONT_DIR = (
    Path(os.environ["WINDIR"]) / "Fonts"
    if os.environ.get("WINDIR")
    else Path("/usr/share/fonts/truetype")
)
FONT_DIR = Path(os.environ.get("CV_FIXTURE_FONT_DIR", DEFAULT_FONT_DIR))
POPPLER_DIR = (
    Path(os.environ["POPPLER_DIR"])
    if os.environ.get("POPPLER_DIR")
    else None
)
QUERY_DATE = "2026-09-29"
SYNTHETIC_NOTICE = "FICTIONAL TEST DATA - NOT A REAL PERSON"


def register_fonts() -> None:
    fonts = {
        "Arial": "arial.ttf",
        "Arial-Bold": "arialbd.ttf",
        "Arial-Italic": "ariali.ttf",
        "Calibri": "calibri.ttf",
        "Calibri-Bold": "calibrib.ttf",
        "Georgia": "georgia.ttf",
        "Georgia-Bold": "georgiab.ttf",
        "Times": "times.ttf",
        "Times-Bold": "timesbd.ttf",
        "Times-Italic": "timesi.ttf",
        "MicrosoftYaHei": "msyh.ttc",
        "MicrosoftYaHei-Bold": "msyhbd.ttc",
        "DengXian": "Deng.ttf",
        "DengXian-Bold": "Dengb.ttf",
    }
    for name, filename in fonts.items():
        path = FONT_DIR / filename
        if path.suffix.lower() == ".ttc":
            pdfmetrics.registerFont(TTFont(name, str(path), subfontIndex=0))
        else:
            pdfmetrics.registerFont(TTFont(name, str(path)))


def add_page_footer(
    c: canvas.Canvas,
    doc: BaseDocTemplate,
    *,
    font: str = "Arial",
    size: float = 7,
    color: colors.Color = colors.HexColor("#777777"),
    left_text: str = "",
) -> None:
    c.saveState()
    c.setFont(font, size)
    c.setFillColor(color)
    if left_text:
        c.drawString(doc.leftMargin, 10 * mm, left_text)
    c.drawRightString(
        A4[0] - doc.rightMargin,
        10 * mm,
        f"Page {doc.page} | {SYNTHETIC_NOTICE}",
    )
    c.restoreState()


def pdf_meta(path: Path, title: str, subject: str) -> None:
    """Set document-level metadata without changing page content."""
    from pypdf import PdfReader, PdfWriter

    reader = PdfReader(str(path))
    writer = PdfWriter()
    for page in reader.pages:
        writer.add_page(page)
    writer.add_metadata(
        {
            "/Title": title,
            "/Author": "Synthetic Test Data",
            "/Subject": subject,
            "/Keywords": "synthetic, fictional, CV, test fixture",
        }
    )
    with path.open("wb") as handle:
        writer.write(handle)


def wrapped_lines(
    text: str, font: str, size: float, width: float
) -> list[str]:
    lines: list[str] = []
    for raw_line in text.splitlines() or [""]:
        current = ""
        for character in raw_line:
            candidate = current + character
            if current and pdfmetrics.stringWidth(candidate, font, size) > width:
                lines.append(current.rstrip())
                current = character.lstrip()
            else:
                current = candidate
        lines.append(current.rstrip())
    return lines


def draw_wrapped(
    c: canvas.Canvas,
    text: str,
    x: float,
    y: float,
    width: float,
    *,
    font: str,
    size: float,
    leading: float,
    color: colors.Color = colors.black,
    max_lines: int | None = None,
) -> float:
    c.setFont(font, size)
    c.setFillColor(color)
    lines = wrapped_lines(text, font, size, width)
    if max_lines is not None:
        lines = lines[:max_lines]
    for line in lines:
        c.drawString(x, y, line)
        y -= leading
    return y


def draw_bullets(
    c: canvas.Canvas,
    bullets: Iterable[str],
    x: float,
    y: float,
    width: float,
    *,
    font: str,
    size: float,
    leading: float,
    color: colors.Color = colors.black,
) -> float:
    for bullet in bullets:
        c.setFont(font, size)
        c.setFillColor(color)
        c.drawString(x, y, "-")
        y = draw_wrapped(
            c,
            bullet,
            x + 9,
            y,
            width - 9,
            font=font,
            size=size,
            leading=leading,
            color=color,
        )
        y -= 2
    return y


def add_footer_to_static_page(c: canvas.Canvas, page_number: int) -> None:
    c.setFont("Arial", 7)
    c.setFillColor(colors.HexColor("#7A7A7A"))
    c.drawRightString(
        A4[0] - 14 * mm,
        9 * mm,
        f"Page {page_number} | {SYNTHETIC_NOTICE}",
    )


def build_academic_pdf(path: Path) -> None:
    body = ParagraphStyle(
        "AcademicBody",
        fontName="Times",
        fontSize=8.7,
        leading=10.4,
        textColor=colors.HexColor("#111111"),
        spaceAfter=1,
    )
    bullet = ParagraphStyle(
        "AcademicBullet",
        parent=body,
        leftIndent=11,
        firstLineIndent=-8,
        spaceAfter=1.5,
    )
    section = ParagraphStyle(
        "AcademicSection",
        fontName="Times-Bold",
        fontSize=10.4,
        leading=11.8,
        textColor=colors.HexColor("#111111"),
        spaceBefore=7,
        spaceAfter=2.5,
    )
    small = ParagraphStyle(
        "AcademicSmall",
        parent=body,
        fontSize=7.9,
        leading=9.2,
        textColor=colors.HexColor("#333333"),
    )

    doc = BaseDocTemplate(
        str(path),
        pagesize=letter,
        leftMargin=16.5 * mm,
        rightMargin=16.5 * mm,
        topMargin=14 * mm,
        bottomMargin=15 * mm,
        title="Evelyn R. Hart - Academic CV Fixture",
        author="Synthetic Test Data",
        subject="Synthetic standard academic CV",
    )
    frame = Frame(
        doc.leftMargin,
        doc.bottomMargin,
        doc.width,
        doc.height,
        id="academic",
    )
    doc.addPageTemplates(
        [
            PageTemplate(
                id="page",
                frames=[frame],
                onPage=lambda c, d: add_page_footer(
                    c,
                    d,
                    font="Times",
                    left_text="Synthetic academic CV",
                ),
            )
        ]
    )

    story: list[object] = []
    story.append(
        Paragraph(
            "EVELYN R. HART, Ph.D.",
            ParagraphStyle(
                "AcademicName",
                fontName="Times-Bold",
                fontSize=20,
                leading=22,
                alignment=TA_CENTER,
                textColor=colors.HexColor("#111111"),
            ),
        )
    )
    story.append(
        Paragraph(
            "Assistant Professor of Human-Centered AI",
            ParagraphStyle(
                "AcademicTitle",
                fontName="Times-Italic",
                fontSize=11,
                leading=13,
                alignment=TA_CENTER,
                textColor=colors.HexColor("#333333"),
                spaceAfter=3,
            ),
        )
    )
    story.append(
        Paragraph(
            "Northbridge Institute of Technology, Department of Computer Science<br/>"
            "evelyn.hart@example.edu | +1 555-0107 | evelyn-hart.example | "
            "github.com/synthetic-ehart",
            ParagraphStyle(
                "AcademicContact",
                fontName="Times",
                fontSize=8.5,
                leading=10,
                alignment=TA_CENTER,
                textColor=colors.HexColor("#333333"),
                spaceAfter=5,
            ),
        )
    )
    story.append(
        HRFlowable(
            width="100%",
            thickness=0.6,
            color=colors.HexColor("#666666"),
            spaceBefore=0,
            spaceAfter=1,
        )
    )

    story.append(Paragraph("RESEARCH PROFILE", section))
    story.append(
        Paragraph(
            "I study how machine-learning systems behave when they encounter changes in "
            "people, workflows, and data. My current work combines model calibration, "
            "causal evaluation, and clinician-facing decision support. I teach machine "
            "learning and lead a small interdisciplinary research group.",
            body,
        )
    )

    story.append(Paragraph("EDUCATION", section))
    education_rows = [
        [
            Paragraph("<b>Ph.D., Computer Science</b><br/>University of Calder", body),
            Paragraph("2015 - 2019<br/>Dissertation: Robust Prediction under Distribution Shift", small),
        ],
        [
            Paragraph("<b>M.Sc., Statistics</b><br/>University of Calder", body),
            Paragraph("2013 - 2015<br/>Thesis: Semi-supervised Risk Estimation", small),
        ],
        [
            Paragraph("<b>B.Sc., Mathematics</b><br/>Westmere College", body),
            Paragraph("2009 - 2013<br/>First-class honours", small),
        ],
    ]
    edu_table = Table(education_rows, colWidths=[118 * mm, 56 * mm], hAlign="LEFT")
    edu_table.setStyle(
        TableStyle(
            [
                ("VALIGN", (0, 0), (-1, -1), "TOP"),
                ("BOTTOMPADDING", (0, 0), (-1, -1), 3),
                ("TOPPADDING", (0, 0), (-1, -1), 1),
                ("LEFTPADDING", (0, 0), (-1, -1), 0),
                ("RIGHTPADDING", (0, 0), (-1, -1), 0),
            ]
        )
    )
    story.append(edu_table)

    story.append(Paragraph("ACADEMIC APPOINTMENTS", section))
    appointments = [
        [
            Paragraph(
                "<b>Assistant Professor</b>, Computer Science<br/>"
                "Northbridge Institute of Technology",
                body,
            ),
            Paragraph("2022 - Present", body),
        ],
        [
            Paragraph(
                "<b>Postdoctoral Research Associate</b>, Health Data Science<br/>"
                "Calder Centre for Clinical AI",
                body,
            ),
            Paragraph("2019 - 2022", body),
        ],
    ]
    appt_table = Table(appointments, colWidths=[132 * mm, 42 * mm])
    appt_table.setStyle(
        TableStyle(
            [
                ("VALIGN", (0, 0), (-1, -1), "TOP"),
                ("ALIGN", (1, 0), (1, -1), "RIGHT"),
                ("BOTTOMPADDING", (0, 0), (-1, -1), 3),
                ("TOPPADDING", (0, 0), (-1, -1), 1),
                ("LEFTPADDING", (0, 0), (-1, -1), 0),
                ("RIGHTPADDING", (0, 0), (-1, -1), 0),
            ]
        )
    )
    story.append(appt_table)

    story.append(Paragraph("SELECTED FUNDING", section))
    for text in [
        "<b>PI</b>, Adaptation-Aware Decision Support for Community Clinics, "
        "Arcadia Health Research Council, 2025 - 2028, GBP 420,000.",
        "<b>PI</b>, Human-in-the-Loop Calibration for Rare Disease Screening, "
        "Northbridge Seed Fund, 2023 - 2025, GBP 75,000.",
        "<b>Co-I</b>, Responsible Deployment of Predictive Models, Calder-Industry "
        "Partnership, 2021 - 2023, GBP 180,000.",
    ]:
        story.append(Paragraph(f"&bull; {text}", bullet))

    story.append(Paragraph("AWARDS AND HONOURS", section))
    for text in [
        "Early Career Research Prize, Calder Society for Computing, 2024.",
        "Best Paper Award, Synthetic AI for Health Conference, 2023.",
        "Outstanding Doctoral Dissertation, University of Calder, 2020.",
    ]:
        story.append(Paragraph(f"&bull; {text}", bullet))

    story.append(Paragraph("ACADEMIC SERVICE", section))
    story.append(
        Paragraph(
            "Reviewer for Responsible Machine Learning Review (2023 - Present) and "
            "Clinical Data Science Forum (2022 - Present). Program committee member, "
            "Synthetic AI for Health Conference (2023 - 2026).",
            body,
        )
    )

    story.append(PageBreak())
    story.append(Paragraph("PUBLICATIONS", section))
    story.append(
        Paragraph(
            "Names are fictional. Citation counts are shown only to exercise extraction "
            "of numeric impact fields.",
            small,
        )
    )
    publications = [
        "Hart, E. R., Imani, N., and Bell, J. (2025). Calibration drift in clinical "
        "decision support after deployment. Journal of Responsible Machine Learning, "
        "12(2), 101-119. Citations: 41.",
        "Hart, E. R., and Patel, S. (2024). When does causal evaluation change model "
        "selection? Proceedings of the Synthetic AI for Health Conference, 88-101. "
        "Citations: 28.",
        "Imani, N., Hart, E. R., and Olsen, M. (2024). Uncertainty communication for "
        "high-stakes predictions. Clinical Data Science Forum, 7(1), 33-49. Citations: 18.",
        "Hart, E. R., Chen, L., and Ward, P. (2023). A benchmark for temporal robustness "
        "in tabular healthcare models. Journal of Applied Machine Learning, 9(4), "
        "410-432. Citations: 64.",
        "Hart, E. R. (2023). Distribution shift is not only a data problem. Responsible "
        "Machine Learning Review, 5(3), 201-218. Citations: 22.",
        "Olsen, M., Hart, E. R., and Imani, N. (2022). Participatory model evaluation "
        "with clinical teams. Proceedings of the Northbridge Symposium on Computing, "
        "144-153. Citations: 31.",
        "Hart, E. R., and Bell, J. (2021). Risk-aware abstention for rare disease "
        "screening. Health Data Science Letters, 4(2), 77-91. Citations: 37.",
        "Patel, S., Hart, E. R., and Ward, P. (2020). Robustness under subgroup shift "
        "with limited labels. Journal of Statistical Learning, 14(1), 5-27. Citations: 52.",
        "Hart, E. R. (2019). Semi-supervised risk estimation under unobserved "
        "confounding. Calder Computing Monographs, 3, 1-38. Citations: 11.",
    ]
    for index, publication in enumerate(publications, 1):
        story.append(
            Paragraph(
                f"<b>{index}.</b> {publication}",
                ParagraphStyle(
                    "Publication",
                    parent=body,
                    leftIndent=10,
                    firstLineIndent=-10,
                    spaceAfter=3,
                ),
            )
        )

    story.append(Paragraph("TEACHING", section))
    story.append(
        Paragraph(
            "CS 421 Machine Learning (2022 - Present, 78-112 students); "
            "CS 615 Evaluation of AI Systems (2023 - Present, graduate). "
            "Guest lecturer, Clinical AI Ethics, Calder Centre for Clinical AI, 2021.",
            body,
        )
    )
    story.append(Paragraph("TECHNICAL SKILLS", section))
    story.append(
        Paragraph(
            "Python, R, SQL, PyTorch, scikit-learn, causal inference, calibration, "
            "data visualization, research supervision, grant writing.",
            body,
        )
    )
    doc.build(story)
    pdf_meta(path, "Evelyn R. Hart - Academic CV Fixture", "Synthetic academic CV")


def build_ats_pdf(path: Path) -> None:
    body = ParagraphStyle(
        "AtsBody",
        fontName="Arial",
        fontSize=8.25,
        leading=9.7,
        textColor=colors.HexColor("#111111"),
        spaceAfter=1.3,
    )
    bullet = ParagraphStyle(
        "AtsBullet",
        parent=body,
        leftIndent=9,
        firstLineIndent=-7,
    )
    heading = ParagraphStyle(
        "AtsHeading",
        fontName="Arial-Bold",
        fontSize=10.4,
        leading=11.2,
        textColor=colors.HexColor("#111111"),
        spaceBefore=6,
        spaceAfter=2.5,
    )
    doc = SimpleDocTemplate(
        str(path),
        pagesize=letter,
        leftMargin=16 * mm,
        rightMargin=16 * mm,
        topMargin=13 * mm,
        bottomMargin=14 * mm,
        title="Marcus T. Chen - ATS Resume Fixture",
        author="Synthetic Test Data",
        subject="Synthetic ATS-friendly industry resume",
    )
    story: list[object] = []
    story.append(
        Paragraph(
            "MARCUS T. CHEN",
            ParagraphStyle(
                "AtsName",
                fontName="Arial-Bold",
                fontSize=18,
                leading=20,
                alignment=TA_LEFT,
                textColor=colors.HexColor("#111111"),
            ),
        )
    )
    story.append(
        Paragraph(
            "Senior Data Scientist | Seattle, WA | marcus.chen@example.com | "
            "+1 555-0118 | linkedin.com/in/synthetic-mchen",
            ParagraphStyle(
                "AtsTitle",
                parent=body,
                fontSize=8.6,
                textColor=colors.HexColor("#333333"),
                spaceAfter=4,
            ),
        )
    )
    story.append(Paragraph("PROFESSIONAL SUMMARY", heading))
    story.append(
        Paragraph(
            "Data scientist with 8 years of experience building forecasting, pricing, "
            "and experimentation systems. Leads projects from problem definition through "
            "deployment and measurement. Strong Python, SQL, cloud, and stakeholder "
            "communication skills.",
            body,
        )
    )
    story.append(Paragraph("CORE SKILLS", heading))
    story.append(
        Paragraph(
            "Python, SQL, PyTorch, scikit-learn, A/B testing, causal inference, "
            "forecasting, Snowflake, AWS, Airflow, Tableau, stakeholder management.",
            body,
        )
    )
    story.append(Paragraph("PROFESSIONAL EXPERIENCE", heading))
    story.append(
        Paragraph(
            "<b>Rainline Analytics, Seattle, WA - Senior Data Scientist</b> "
            "(2021 - Present)",
            body,
        )
    )
    for text in [
        "Led a demand forecasting program across 140 stores and improved 8-week "
        "forecast accuracy by 17% through hierarchical models and promotion features.",
        "Designed 23 controlled experiments and established an analysis review that "
        "reduced repeated metric disputes by 40%.",
        "Deployed a pricing sensitivity service with engineering partners; it processes "
        "1.2 million daily events and supports weekly commercial decisions.",
    ]:
        story.append(Paragraph(f"&bull; {text}", bullet))
    story.append(
        Paragraph(
            "<b>Cascade Retail Group, Portland, OR - Data Scientist</b> "
            "(2018 - 2021)",
            body,
        )
    )
    for text in [
        "Built churn models covering 620,000 customers and identified retention actions "
        "that increased 90-day repeat purchase by 6.4%.",
        "Automated a reporting suite that replaced 30 hours of manual analysis per month.",
    ]:
        story.append(Paragraph(f"&bull; {text}", bullet))
    story.append(Paragraph("EDUCATION", heading))
    story.append(
        Paragraph(
            "<b>M.S., Statistics</b>, Lakeview University, 2018<br/>"
            "<b>B.S., Economics</b>, Lakeview University, 2016",
            body,
        )
    )
    story.append(Paragraph("SELECTED PROJECTS", heading))
    story.append(
        Paragraph(
            "<b>Open Demand Simulator</b> - Maintained a synthetic Python library for "
            "testing hierarchical forecasting workflows; 420 GitHub stars (sample data).",
            bullet,
        )
    )
    story.append(
        Paragraph(
            "<b>Retail Experiment Library</b> - Created reusable statistical notebooks "
            "adopted by four analytics teams.",
            bullet,
        )
    )
    story.append(Paragraph("CERTIFICATIONS", heading))
    story.append(
        Paragraph(
            "AWS Certified Machine Learning - Specialty (2024); "
            "Professional Certificate in Causal Inference (2022).",
            body,
        )
    )
    doc.build(story)
    pdf_meta(path, "Marcus T. Chen - ATS Resume Fixture", "Synthetic ATS resume")


def build_two_column_pdf(path: Path) -> None:
    c = canvas.Canvas(str(path), pagesize=A4)
    c.setTitle("Priya N. Rao - Two Column Resume Fixture")
    c.setAuthor("Synthetic Test Data")
    c.setSubject("Synthetic modern two-column resume")
    sidebar_width = 68 * mm
    page_w, page_h = A4

    def draw_page_base(page_number: int) -> None:
        c.setFillColor(colors.HexColor("#17324D"))
        c.rect(0, 0, sidebar_width, page_h, fill=1, stroke=0)
        c.setFillColor(colors.HexColor("#F3F6F8"))
        c.rect(sidebar_width, 0, page_w - sidebar_width, page_h, fill=1, stroke=0)
        add_footer_to_static_page(c, page_number)

    draw_page_base(1)
    c.setFillColor(colors.white)
    c.setFont("Arial-Bold", 24)
    c.drawString(13 * mm, page_h - 24 * mm, "PRIYA")
    c.drawString(13 * mm, page_h - 31 * mm, "N. RAO")
    c.setFont("Arial", 11)
    c.setFillColor(colors.HexColor("#B8D6E8"))
    c.drawString(13 * mm, page_h - 39 * mm, "Product Analytics Lead")

    y = page_h - 52 * mm
    c.setFillColor(colors.HexColor("#8FC7E8"))
    c.setFont("Arial-Bold", 9)
    c.drawString(13 * mm, y, "CONTACT")
    y -= 6 * mm
    for line in [
        "priya.rao@example.com",
        "+1 555-0134",
        "Seattle, WA",
        "linkedin.com/in/synthetic-prao",
        "priya-rao.example",
    ]:
        y = draw_wrapped(
            c,
            line,
            13 * mm,
            y,
            sidebar_width - 24 * mm,
            font="Arial",
            size=8,
            leading=4.5 * mm,
            color=colors.white,
        )
    y -= 5 * mm
    c.setFillColor(colors.HexColor("#8FC7E8"))
    c.setFont("Arial-Bold", 9)
    c.drawString(13 * mm, y, "CORE SKILLS")
    y -= 6 * mm
    for skill in [
        "Experimentation",
        "Causal inference",
        "Product metrics",
        "SQL and Python",
        "Forecast planning",
        "Stakeholder facilitation",
    ]:
        c.setFillColor(colors.HexColor("#8FC7E8"))
        c.circle(14 * mm, y + 1.5 * mm, 1.2 * mm, fill=1, stroke=0)
        y = draw_wrapped(
            c,
            skill,
            17 * mm,
            y,
            sidebar_width - 28 * mm,
            font="Arial",
            size=8,
            leading=4.5 * mm,
            color=colors.white,
        )
    y -= 4 * mm
    c.setFillColor(colors.HexColor("#8FC7E8"))
    c.setFont("Arial-Bold", 9)
    c.drawString(13 * mm, y, "EDUCATION")
    y -= 6 * mm
    y = draw_wrapped(
        c,
        "M.S. Applied Statistics\nLakeview University\n2016 - 2018",
        13 * mm,
        y,
        sidebar_width - 24 * mm,
        font="Arial",
        size=8,
        leading=4.5 * mm,
        color=colors.white,
    )
    y -= 4 * mm
    y = draw_wrapped(
        c,
        "B.A. Economics\nWestmere College\n2012 - 2016",
        13 * mm,
        y,
        sidebar_width - 24 * mm,
        font="Arial",
        size=8,
        leading=4.5 * mm,
        color=colors.white,
    )
    y -= 4 * mm
    c.setFillColor(colors.HexColor("#8FC7E8"))
    c.setFont("Arial-Bold", 9)
    c.drawString(13 * mm, y, "LANGUAGES")
    y -= 6 * mm
    draw_wrapped(
        c,
        "English (native)\nSpanish (professional)\nHindi (conversational)",
        13 * mm,
        y,
        sidebar_width - 24 * mm,
        font="Arial",
        size=8,
        leading=4.5 * mm,
        color=colors.white,
    )

    x = sidebar_width + 14 * mm
    right_width = page_w - x - 14 * mm
    y = page_h - 22 * mm
    c.setFillColor(colors.HexColor("#17324D"))
    c.setFont("Georgia-Bold", 15)
    c.drawString(x, y, "PROFILE")
    y -= 8 * mm
    y = draw_wrapped(
        c,
        "Product analytics leader who turns ambiguous customer questions into "
        "measurement systems, experiments, and decisions teams can act on. "
        "Nine years across marketplace, subscription, and retail products.",
        x,
        y,
        right_width,
        font="Arial",
        size=9,
        leading=5.2 * mm,
        color=colors.HexColor("#27313A"),
    )
    y -= 6 * mm
    c.setFillColor(colors.HexColor("#17324D"))
    c.setFont("Georgia-Bold", 15)
    c.drawString(x, y, "EXPERIENCE")
    y -= 9 * mm

    roles = [
        (
            "Product Analytics Lead",
            "Meridian Market | Seattle, WA | 2022 - Present",
            [
                "Built a shared metric layer for 6 product teams and reduced conflicting "
                "weekly reporting from 12 definitions to 1 approved metric set.",
                "Designed a seller-quality experiment that lifted repeat listings by 9.3% "
                "without increasing marketplace abuse reports.",
            ],
        ),
        (
            "Senior Product Analyst",
            "Harborline Subscription | Seattle, WA | 2019 - 2022",
            [
                "Led retention diagnostics across 4.1 million accounts and identified "
                "two onboarding changes that improved 90-day retention by 5.8%.",
                "Created a causal readout for a pricing test spanning 18 markets.",
            ],
        ),
    ]
    for role, meta, bullets in roles:
        c.setFillColor(colors.HexColor("#27313A"))
        c.setFont("Arial-Bold", 10)
        c.drawString(x, y, role)
        y -= 4.5 * mm
        c.setFillColor(colors.HexColor("#607080"))
        c.setFont("Arial-Italic", 8)
        c.drawString(x, y, meta)
        y -= 5.2 * mm
        y = draw_bullets(
            c,
            bullets,
            x,
            y,
            right_width,
            font="Arial",
            size=8.2,
            leading=4.6 * mm,
            color=colors.HexColor("#27313A"),
        )
        y -= 3 * mm
    c.showPage()

    draw_page_base(2)
    x = sidebar_width + 14 * mm
    y = page_h - 24 * mm
    c.setFillColor(colors.HexColor("#17324D"))
    c.setFont("Georgia-Bold", 15)
    c.drawString(x, y, "SELECTED PRODUCT WINS")
    y -= 10 * mm
    wins = [
        (
            "Forecast Confidence",
            "Introduced prediction intervals into planning tools. Planners reduced "
            "manual overrides by 22% over one quarter.",
        ),
        (
            "Experiment Review",
            "Created a lightweight pre-registration template used by 43 product "
            "experiments and now part of the release checklist.",
        ),
        (
            "Self-Service Analytics",
            "Shipped documented data products that reduced recurring analyst requests "
            "by roughly 30%.",
        ),
    ]
    for title, description in wins:
        y = draw_wrapped(
            c,
            title,
            x,
            y,
            right_width,
            font="Arial-Bold",
            size=10,
            leading=5.2 * mm,
            color=colors.HexColor("#17324D"),
        )
        y -= 1 * mm
        y = draw_wrapped(
            c,
            description,
            x,
            y,
            right_width,
            font="Arial",
            size=8.8,
            leading=5 * mm,
            color=colors.HexColor("#27313A"),
        )
        y -= 6 * mm
    c.setFillColor(colors.HexColor("#17324D"))
    c.setFont("Georgia-Bold", 15)
    c.drawString(x, y, "TOOLS")
    y -= 8 * mm
    y = draw_wrapped(
        c,
        "SQL, Python, R, dbt, Snowflake, Looker, Amplitude, Optimizely, Git, "
        "experiment design, causal graphs, metric governance.",
        x,
        y,
        right_width,
        font="Arial",
        size=9,
        leading=5 * mm,
        color=colors.HexColor("#27313A"),
    )
    c.save()
    pdf_meta(path, "Priya N. Rao - Two Column Resume Fixture", "Synthetic two-column resume")


def build_chinese_campus_pdf(path: Path) -> None:
    c = canvas.Canvas(str(path), pagesize=A4)
    c.setTitle("林雨澄 - 中文校招简历样例")
    c.setAuthor("Synthetic Test Data")
    c.setSubject("Synthetic Chinese campus recruitment CV")
    page_w, page_h = A4
    header_h = 43 * mm

    def page_base(page_number: int) -> None:
        c.setFillColor(colors.HexColor("#FBFCFE"))
        c.rect(0, 0, page_w, page_h, fill=1, stroke=0)
        c.setFillColor(colors.HexColor("#1D5C8F"))
        c.rect(0, page_h - header_h, page_w, header_h, fill=1, stroke=0)
        c.setFillColor(colors.HexColor("#53A7D8"))
        c.rect(0, page_h - header_h - 2 * mm, page_w, 2 * mm, fill=1, stroke=0)
        add_footer_to_static_page(c, page_number)

    page_base(1)
    c.setFillColor(colors.white)
    c.setFont("MicrosoftYaHei-Bold", 23)
    c.drawString(18 * mm, page_h - 21 * mm, "林雨澄")
    c.setFont("MicrosoftYaHei", 10)
    c.drawString(18 * mm, page_h - 29 * mm, "求职方向：数据分析师 / AI 产品运营")
    c.setFont("MicrosoftYaHei", 8.5)
    c.drawString(
        18 * mm,
        page_h - 36 * mm,
        "电话：138-0000-0000   |   邮箱：yucheng.lin@example.com   |   地点：上海",
    )
    c.setFillColor(colors.white)
    c.roundRect(page_w - 43 * mm, page_h - 32 * mm, 25 * mm, 22 * mm, 2 * mm, fill=0, stroke=1)
    c.setFont("MicrosoftYaHei", 7)
    c.drawCentredString(page_w - 30.5 * mm, page_h - 21 * mm, "照片占位")

    x = 18 * mm
    y = page_h - 54 * mm
    right = page_w - 18 * mm
    content_w = right - x

    def section_title(title: str) -> None:
        nonlocal y
        c.setFillColor(colors.HexColor("#1D5C8F"))
        c.rect(x, y - 1.5 * mm, 3 * mm, 6 * mm, fill=1, stroke=0)
        c.setFont("MicrosoftYaHei-Bold", 12)
        c.drawString(x + 6 * mm, y, title)
        y -= 8 * mm

    section_title("教育背景")
    c.setFont("MicrosoftYaHei-Bold", 9.5)
    c.setFillColor(colors.HexColor("#222222"))
    c.drawString(x, y, "星海大学")
    c.setFont("MicrosoftYaHei", 8.5)
    c.drawRightString(right, y, "2022.09 - 2026.06")
    y -= 4.8 * mm
    y = draw_wrapped(
        c,
        "数据科学与大数据技术 本科 | GPA 3.8/4.0 | 专业排名 5/126",
        x,
        y,
        content_w,
        font="MicrosoftYaHei",
        size=8.5,
        leading=4.8 * mm,
        color=colors.HexColor("#444444"),
    )
    y = draw_wrapped(
        c,
        "核心课程：机器学习、数据库系统、概率统计、数据可视化、用户研究",
        x,
        y,
        content_w,
        font="MicrosoftYaHei",
        size=8.3,
        leading=4.7 * mm,
        color=colors.HexColor("#555555"),
    )
    y -= 3 * mm

    section_title("实习经历")
    internships = [
        (
            "云杉科技 | 数据分析实习生",
            "2025.07 - 2025.11",
            [
                "搭建用户留存看板，覆盖 6 条产品线，将周报制作时间从 5 小时缩短到 40 分钟。",
                "使用 SQL 与 Python 分析新用户激活漏斗，提出 3 项改版建议，其中 2 项进入 A/B 测试。",
            ],
        ),
        (
            "星环零售 | 增长运营实习生",
            "2024.07 - 2024.10",
            [
                "负责会员活动的数据复盘，按渠道与城市拆分转化率，输出 12 页分析报告。",
                "联合设计问卷并访谈 18 位用户，整理 46 条需求标签。",
            ],
        ),
    ]
    for role, dates, bullets in internships:
        c.setFont("MicrosoftYaHei-Bold", 9.2)
        c.setFillColor(colors.HexColor("#222222"))
        c.drawString(x, y, role)
        c.setFont("MicrosoftYaHei", 8.2)
        c.drawRightString(right, y, dates)
        y -= 5 * mm
        y = draw_bullets(
            c,
            bullets,
            x,
            y,
            content_w,
            font="MicrosoftYaHei",
            size=8.3,
            leading=4.8 * mm,
            color=colors.HexColor("#444444"),
        )
        y -= 2.5 * mm

    section_title("项目经历")
    projects = [
        (
            "校园二手交易需求预测 | 课程项目负责人",
            "2025.03 - 2025.06",
            "清洗 3.2 万条交易记录，构建随机森林与时间序列基线，将节假日销量预测误差降低 18%。",
        ),
        (
            "智能简历信息抽取原型 | 个人项目",
            "2024.12 - 2025.02",
            "使用 Python 与正则规则提取教育、技能和时间信息，为 120 份匿名样本生成结构化字段。",
        ),
    ]
    for title, dates, description in projects:
        c.setFont("MicrosoftYaHei-Bold", 9.2)
        c.setFillColor(colors.HexColor("#222222"))
        c.drawString(x, y, title)
        c.setFont("MicrosoftYaHei", 8.2)
        c.drawRightString(right, y, dates)
        y -= 5 * mm
        y = draw_wrapped(
            c,
            description,
            x,
            y,
            content_w,
            font="MicrosoftYaHei",
            size=8.3,
            leading=4.8 * mm,
            color=colors.HexColor("#444444"),
        )
        y -= 3 * mm

    c.showPage()
    page_base(2)
    x = 18 * mm
    y = page_h - 22 * mm

    section_title("专业技能")
    skills = [
        ("数据分析", "SQL、Python、Pandas、NumPy、数据清洗、漏斗分析"),
        ("机器学习", "scikit-learn、基础模型评估、A/B 测试"),
        ("可视化", "Tableau、Power BI、Matplotlib"),
        ("协作与语言", "中文母语、英语 CET-6、跨团队沟通"),
    ]
    for label, value in skills:
        c.setFont("MicrosoftYaHei-Bold", 8.8)
        c.setFillColor(colors.HexColor("#1D5C8F"))
        c.drawString(x, y, label)
        draw_wrapped(
            c,
            value,
            x + 27 * mm,
            y,
            content_w - 27 * mm,
            font="MicrosoftYaHei",
            size=8.3,
            leading=4.8 * mm,
            color=colors.HexColor("#333333"),
        )
        y -= 9 * mm
    y -= 2 * mm

    section_title("获奖与证书")
    awards = [
        "2025 全国大学生数据建模挑战赛 一等奖（虚构赛事）",
        "2024 星海大学优秀学生奖学金 二等奖",
        "Tableau Desktop Specialist（2025）",
    ]
    y = draw_bullets(
        c,
        awards,
        x,
        y,
        content_w,
        font="MicrosoftYaHei",
        size=8.4,
        leading=5 * mm,
        color=colors.HexColor("#333333"),
    )
    y -= 2 * mm

    section_title("校园经历")
    y = draw_bullets(
        c,
        [
            "数据科学协会活动负责人，组织 7 场 Python 工作坊，累计参与 260 人次。",
            "院学生会宣传部成员，负责活动数据统计与复盘。",
        ],
        x,
        y,
        content_w,
        font="MicrosoftYaHei",
        size=8.4,
        leading=5 * mm,
        color=colors.HexColor("#333333"),
    )
    y -= 3 * mm

    section_title("自我评价")
    draw_wrapped(
        c,
        "习惯先确认问题，再选择指标和方法。能够独立完成取数、分析、结论汇报，"
        "也愿意主动与产品、运营和工程同学对齐口径。",
        x,
        y,
        content_w,
        font="MicrosoftYaHei",
        size=8.6,
        leading=5.3 * mm,
        color=colors.HexColor("#333333"),
    )
    c.save()
    pdf_meta(path, "林雨澄 - 中文校招简历样例", "Synthetic Chinese campus CV")


def build_infographic_pdf(path: Path) -> None:
    c = canvas.Canvas(str(path), pagesize=A4)
    c.setTitle("Sofia Alvarez - Infographic Resume Fixture")
    c.setAuthor("Synthetic Test Data")
    c.setSubject("Synthetic creative infographic resume")
    page_w, page_h = A4
    c.setFillColor(colors.HexColor("#F6F4EF"))
    c.rect(0, 0, page_w, page_h, fill=1, stroke=0)
    c.setFillColor(colors.HexColor("#2D4A43"))
    c.rect(0, page_h - 54 * mm, page_w, 54 * mm, fill=1, stroke=0)
    c.setFillColor(colors.HexColor("#E7A64A"))
    c.circle(page_w - 25 * mm, page_h - 34 * mm, 23 * mm, fill=1, stroke=0)
    c.setFillColor(colors.white)
    c.setFont("Georgia-Bold", 26)
    c.drawString(18 * mm, page_h - 24 * mm, "SOFIA ALVAREZ")
    c.setFont("Arial", 11)
    c.drawString(18 * mm, page_h - 33 * mm, "UX RESEARCHER / SERVICE DESIGNER")
    c.setFont("Arial", 8.2)
    c.drawString(
        18 * mm,
        page_h - 43 * mm,
        "sofia.alvarez@example.com  |  +34 555 0190  |  Madrid  |  sofia-alvarez.example",
    )

    left_x = 18 * mm
    y = page_h - 67 * mm
    c.setFillColor(colors.HexColor("#2D4A43"))
    c.setFont("Georgia-Bold", 14)
    c.drawString(left_x, y, "PROFILE")
    y -= 8 * mm
    y = draw_wrapped(
        c,
        "Researcher who connects field evidence, service operations, and product "
        "decisions. I plan studies, synthesize behavior, and help teams agree on what "
        "to test next.",
        left_x,
        y,
        80 * mm,
        font="Arial",
        size=9,
        leading=5.1 * mm,
        color=colors.HexColor("#2C3330"),
    )
    y -= 6 * mm
    c.setFillColor(colors.HexColor("#2D4A43"))
    c.setFont("Georgia-Bold", 14)
    c.drawString(left_x, y, "TOOLBOX")
    y -= 7 * mm
    tool_y = y
    tools = [
        ("Field research", 0.92, "#2D4A43"),
        ("Usability testing", 0.88, "#3E6B60"),
        ("Survey design", 0.78, "#5F8D80"),
        ("Journey mapping", 0.84, "#E7A64A"),
        ("Service blueprints", 0.72, "#D4774E"),
    ]
    for index, (tool, score, color) in enumerate(tools):
        bar_y = tool_y - index * 10 * mm
        c.setFillColor(colors.HexColor("#D9DEDB"))
        c.roundRect(left_x, bar_y, 58 * mm, 3 * mm, 1.5 * mm, fill=1, stroke=0)
        c.setFillColor(colors.HexColor(color))
        c.roundRect(left_x, bar_y, 58 * mm * score, 3 * mm, 1.5 * mm, fill=1, stroke=0)
        c.setFillColor(colors.HexColor("#2C3330"))
        c.setFont("Arial", 7.5)
        c.drawString(left_x, bar_y + 4.2 * mm, tool)
        c.drawRightString(left_x + 58 * mm, bar_y + 4.2 * mm, f"{int(score * 100)}%")
    y = tool_y - len(tools) * 10 * mm - 2 * mm

    c.setFillColor(colors.HexColor("#2D4A43"))
    c.setFont("Georgia-Bold", 14)
    c.drawString(left_x, y, "EDUCATION")
    y -= 7 * mm
    y = draw_wrapped(
        c,
        "M.A. Service Design\nMadrid Design School\n2017 - 2019",
        left_x,
        y,
        74 * mm,
        font="Arial",
        size=8.6,
        leading=4.8 * mm,
        color=colors.HexColor("#2C3330"),
    )
    y -= 4 * mm
    draw_wrapped(
        c,
        "B.A. Sociology\nUniversity of Valencia\n2013 - 2017",
        left_x,
        y,
        74 * mm,
        font="Arial",
        size=8.6,
        leading=4.8 * mm,
        color=colors.HexColor("#2C3330"),
    )

    line_x = 112 * mm
    top_y = page_h - 68 * mm
    c.setStrokeColor(colors.HexColor("#8EA59E"))
    c.setLineWidth(1.2)
    c.line(line_x, 54 * mm, line_x, top_y + 3 * mm)
    c.setFillColor(colors.HexColor("#2D4A43"))
    c.setFont("Georgia-Bold", 14)
    c.drawString(line_x + 10 * mm, top_y, "EXPERIENCE")

    experiences = [
        (
            "2019 - Present",
            "Senior UX Researcher",
            "Northline Mobility",
            "Led 34 studies across ticketing and accessibility. Built a research "
            "repository used by 7 product squads.",
        ),
        (
            "2018 - 2019",
            "Service Design Consultant",
            "Public Service Lab",
            "Mapped appointment journeys with 6 municipal teams and tested a "
            "low-bandwidth reminder service.",
        ),
        (
            "2017 - 2018",
            "Research Assistant",
            "Madrid Design School",
            "Studied trust in automated transport and coordinated 12 interviews.",
        ),
    ]
    event_y = top_y - 13 * mm
    for dates, role, company, description in experiences:
        c.setFillColor(colors.HexColor("#E7A64A"))
        c.circle(line_x, event_y + 1.5 * mm, 2.5 * mm, fill=1, stroke=0)
        c.setFillColor(colors.HexColor("#6A746F"))
        c.setFont("Arial-Bold", 7.5)
        c.drawString(line_x + 10 * mm, event_y + 5 * mm, dates)
        c.setFillColor(colors.HexColor("#2D4A43"))
        c.setFont("Arial-Bold", 10.5)
        c.drawString(line_x + 10 * mm, event_y - 1.2 * mm, role)
        c.setFont("Arial-Italic", 8.5)
        c.drawString(line_x + 10 * mm, event_y - 6.2 * mm, company)
        desc_y = draw_wrapped(
            c,
            description,
            line_x + 10 * mm,
            event_y - 12 * mm,
            66 * mm,
            font="Arial",
            size=8.2,
            leading=4.6 * mm,
            color=colors.HexColor("#3B4541"),
        )
        event_y = desc_y - 8 * mm

    box_y = 29 * mm
    c.setFillColor(colors.HexColor("#E7A64A"))
    c.roundRect(
        left_x,
        box_y,
        page_w - 36 * mm,
        24 * mm,
        3 * mm,
        fill=1,
        stroke=0,
    )
    c.setFillColor(colors.HexColor("#2D4A43"))
    c.setFont("Georgia-Bold", 13)
    c.drawString(left_x + 8 * mm, box_y + 14 * mm, "LANGUAGES")
    c.setFont("Arial", 8.5)
    c.drawString(
        left_x + 48 * mm,
        box_y + 14 * mm,
        "Spanish native  |  English fluent  |  Portuguese intermediate",
    )
    c.drawString(
        left_x + 8 * mm,
        box_y + 6.5 * mm,
        "Selected talks: Designing Accessible Transit Services; Research with Community Partners.",
    )
    add_footer_to_static_page(c, 1)
    c.save()
    pdf_meta(path, "Sofia Alvarez - Infographic Resume Fixture", "Synthetic infographic resume")


def set_cell_shading(cell, fill: str) -> None:
    tc_pr = cell._tc.get_or_add_tcPr()
    shading = tc_pr.find(qn("w:shd"))
    if shading is None:
        shading = OxmlElement("w:shd")
        tc_pr.append(shading)
    shading.set(qn("w:fill"), fill)


def set_cell_border(cell, color: str = "D9D9D9", size: str = "6") -> None:
    tc_pr = cell._tc.get_or_add_tcPr()
    borders = tc_pr.first_child_found_in("w:tcBorders")
    if borders is None:
        borders = OxmlElement("w:tcBorders")
        tc_pr.append(borders)
    for edge in ("top", "left", "bottom", "right", "insideH", "insideV"):
        element = borders.find(qn(f"w:{edge}"))
        if element is None:
            element = OxmlElement(f"w:{edge}")
            borders.append(element)
        element.set(qn("w:val"), "single")
        element.set(qn("w:sz"), size)
        element.set(qn("w:space"), "0")
        element.set(qn("w:color"), color)


def set_cell_margins(cell, top: int = 90, start: int = 100, bottom: int = 90, end: int = 100) -> None:
    tc = cell._tc
    tc_pr = tc.get_or_add_tcPr()
    tc_mar = tc_pr.first_child_found_in("w:tcMar")
    if tc_mar is None:
        tc_mar = OxmlElement("w:tcMar")
        tc_pr.append(tc_mar)
    for edge, value in (("top", top), ("start", start), ("bottom", bottom), ("end", end)):
        node = tc_mar.find(qn(f"w:{edge}"))
        if node is None:
            node = OxmlElement(f"w:{edge}")
            tc_mar.append(node)
        node.set(qn("w:w"), str(value))
        node.set(qn("w:type"), "dxa")


def set_run_font(run, name: str, size: float, bold: bool = False, color: str = "000000") -> None:
    run.font.name = name
    run.font.size = Pt(size)
    run.font.bold = bold
    run.font.color.rgb = RGBColor.from_string(color)
    run._element.get_or_add_rPr().rFonts.set(qn("w:ascii"), name)
    run._element.get_or_add_rPr().rFonts.set(qn("w:hAnsi"), name)
    run._element.get_or_add_rPr().rFonts.set(qn("w:eastAsia"), name)


def set_repeat_table_header(row) -> None:
    row_props = row._tr.get_or_add_trPr()
    header = OxmlElement("w:tblHeader")
    header.set(qn("w:val"), "true")
    row_props.append(header)


def docx_footer(doc: Document) -> None:
    footer = doc.sections[0].footer
    paragraph = footer.paragraphs[0]
    paragraph.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    run = paragraph.add_run(SYNTHETIC_NOTICE)
    set_run_font(run, "Arial", 7, color="777777")


def set_doc_defaults(doc: Document, normal_font: str, normal_size: float) -> None:
    normal = doc.styles["Normal"]
    normal.font.name = normal_font
    normal.font.size = Pt(normal_size)
    normal._element.rPr.rFonts.set(qn("w:ascii"), normal_font)
    normal._element.rPr.rFonts.set(qn("w:hAnsi"), normal_font)
    normal._element.rPr.rFonts.set(qn("w:eastAsia"), normal_font)
    for section in doc.sections:
        section.page_width = Cm(21)
        section.page_height = Cm(29.7)
        section.top_margin = Cm(1.25)
        section.bottom_margin = Cm(1.15)
        section.left_margin = Cm(1.5)
        section.right_margin = Cm(1.5)
    docx_footer(doc)


def add_para(
    doc: Document,
    text: str = "",
    *,
    font: str = "Calibri",
    size: float = 9,
    bold: bool = False,
    color: str = "000000",
    align: WD_ALIGN_PARAGRAPH | None = None,
    before: float = 0,
    after: float = 0,
    line_spacing: float = 1.0,
) -> None:
    paragraph = doc.add_paragraph()
    if align is not None:
        paragraph.alignment = align
    paragraph.paragraph_format.space_before = Pt(before)
    paragraph.paragraph_format.space_after = Pt(after)
    paragraph.paragraph_format.line_spacing = line_spacing
    run = paragraph.add_run(text)
    set_run_font(run, font, size, bold, color)


def add_docx_section(doc: Document, title: str, color: str = "17324D") -> None:
    add_para(
        doc,
        title.upper(),
        font="Calibri",
        size=11,
        bold=True,
        color=color,
        before=5,
        after=2,
    )


def style_table(table, header_row: int | None = None, accent: str = "D9EAF5") -> None:
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    table.autofit = False
    for row_index, row in enumerate(table.rows):
        for cell in row.cells:
            cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
            set_cell_border(cell)
            set_cell_margins(cell)
            if header_row is not None and row_index == header_row:
                set_cell_shading(cell, accent)


def add_two_col_section(
    doc: Document,
    label: str,
    content_rows: Sequence[tuple[str, str]],
    *,
    label_width: float = 3.5,
    content_width: float = 13.0,
) -> None:
    table = doc.add_table(rows=len(content_rows), cols=2)
    table.style = "Table Grid"
    for index, (heading, body) in enumerate(content_rows):
        label_cell, body_cell = table.rows[index].cells
        label_cell.width = Cm(label_width)
        body_cell.width = Cm(content_width)
        set_cell_shading(label_cell, "DCEAF5")
        label_paragraph = label_cell.paragraphs[0]
        label_run = label_paragraph.add_run(label if index == 0 else "")
        set_run_font(label_run, "Calibri", 9, True, "17324D")
        body_paragraph = body_cell.paragraphs[0]
        body_paragraph.paragraph_format.space_after = Pt(0)
        title_run = body_paragraph.add_run(heading)
        set_run_font(title_run, "Calibri", 9.2, True)
        if body:
            body_paragraph.add_run("\n")
            body_run = body_paragraph.add_run(body)
            set_run_font(body_run, "Calibri", 8.7, False, "333333")
    style_table(table)
    add_para(doc, "", size=2, after=0)


def build_europass_docx(path: Path) -> None:
    doc = Document()
    set_doc_defaults(doc, "Calibri", 9)
    doc.core_properties.title = "Lucas M. Dubois - Europass Style CV Fixture"
    doc.core_properties.author = "Synthetic Test Data"
    doc.core_properties.comments = "Fictional CV fixture for parser testing."

    header = doc.add_table(rows=1, cols=2)
    header.style = "Table Grid"
    left, right = header.rows[0].cells
    left.width = Cm(11.5)
    right.width = Cm(5)
    set_cell_shading(left, "1D5C8F")
    set_cell_shading(right, "D9EAF5")
    left.paragraphs[0].alignment = WD_ALIGN_PARAGRAPH.LEFT
    name_run = left.paragraphs[0].add_run("LUCAS M. DUBOIS")
    set_run_font(name_run, "Calibri", 18, True, "FFFFFF")
    left.paragraphs[0].add_run("\n")
    role_run = left.paragraphs[0].add_run("Data Engineer | Cloud Platforms and Analytics")
    set_run_font(role_run, "Calibri", 10, False, "FFFFFF")
    right.paragraphs[0].alignment = WD_ALIGN_PARAGRAPH.CENTER
    photo_run = right.paragraphs[0].add_run("PHOTO\nPLACEHOLDER")
    set_run_font(photo_run, "Calibri", 9, True, "1D5C8F")
    style_table(header)

    add_para(
        doc,
        "lucas.dubois@example.com | +33 555 0122 | Lyon, France | "
        "linkedin.com/in/synthetic-ldubois",
        size=8.5,
        color="444444",
        before=4,
        after=4,
    )

    add_docx_section(doc, "Personal statement", "1D5C8F")
    add_para(
        doc,
        "Data engineer with seven years of experience building reliable batch and "
        "streaming pipelines for analytics and machine-learning products. I work across "
        "data modeling, orchestration, cloud infrastructure, and data quality.",
        size=8.8,
        color="333333",
        after=3,
    )
    add_two_col_section(
        doc,
        "WORK EXPERIENCE",
        [
            (
                "Senior Data Engineer, Cobalt Mobility | 2022 - Present",
                "Designed event ingestion for 4.8 million daily trips. Reduced pipeline "
                "failure recovery time from 90 minutes to 12 minutes. Mentored three engineers.",
            ),
            (
                "Data Engineer, Alpine Insights | 2019 - 2022",
                "Built an analytics warehouse serving 18 dashboards and 90 users. "
                "Introduced contract tests and release checks for 62 critical tables.",
            ),
        ],
    )
    add_two_col_section(
        doc,
        "EDUCATION AND TRAINING",
        [
            (
                "M.Sc. Data Science, Lyon Technical University | 2017 - 2019",
                "Distributed systems, statistics, data mining. Thesis on streaming data quality.",
            ),
            (
                "B.Sc. Computer Science, University of Grenoble | 2014 - 2017",
                "Graduated with distinction. Exchange semester in Montreal.",
            ),
        ],
    )

    add_docx_section(doc, "LANGUAGE SKILLS", "1D5C8F")
    language_table = doc.add_table(rows=4, cols=3)
    language_table.style = "Table Grid"
    language_data = [
        ("Language", "Understanding", "Speaking / Writing"),
        ("French", "Native", "Native"),
        ("English", "C1", "C1"),
        ("German", "B1", "A2"),
    ]
    for row_index, values in enumerate(language_data):
        for col_index, value in enumerate(values):
            cell = language_table.rows[row_index].cells[col_index]
            run = cell.paragraphs[0].add_run(value)
            set_run_font(run, "Calibri", 8.8, row_index == 0, "FFFFFF" if row_index == 0 else "000000")
            if row_index == 0:
                set_cell_shading(cell, "1D5C8F")
    style_table(language_table, header_row=0, accent="1D5C8F")
    add_para(doc, "", size=2, after=0)

    add_docx_section(doc, "DIGITAL SKILLS", "1D5C8F")
    add_para(
        doc,
        "Python, SQL, Scala, Airflow, dbt, Kafka, Spark, Terraform, AWS, GCP, "
        "Docker, Kubernetes, GitLab CI, observability, data modeling.",
        size=8.8,
        color="333333",
        after=3,
    )
    add_docx_section(doc, "PROJECTS", "1D5C8F")
    projects = [
        (
            "Open Transit Quality Monitor | 2024",
            "Created a synthetic benchmarking toolkit for late-event handling and "
            "published documentation for maintainers.",
        ),
        (
            "Warehouse Cost Explorer | 2023",
            "Built a dashboard that attributed cloud warehouse cost to teams and "
            "query patterns; helped reduce monthly spend by 19%.",
        ),
    ]
    for heading, description in projects:
        add_para(doc, heading, size=9.2, bold=True, color="17324D", before=2, after=1)
        add_para(doc, description, size=8.6, color="333333", after=2)

    add_docx_section(doc, "ADDITIONAL INFORMATION", "1D5C8F")
    add_para(
        doc,
        "Volunteer mentor, Lyon Data Community (2021 - Present). "
        "Speaker, Synthetic Data Engineering Forum (2024).",
        size=8.6,
        color="333333",
    )
    doc.save(path)


def build_german_docx(path: Path) -> None:
    doc = Document()
    set_doc_defaults(doc, "Calibri", 9)
    doc.core_properties.title = "Hannah Keller - German Lebenslauf Fixture"
    doc.core_properties.author = "Synthetic Test Data"

    add_para(
        doc,
        "LEBENSLAUF",
        font="Calibri",
        size=19,
        bold=True,
        color="243447",
        align=WD_ALIGN_PARAGRAPH.CENTER,
        after=8,
    )
    top = doc.add_table(rows=1, cols=2)
    top.style = "Table Grid"
    details, photo = top.rows[0].cells
    details.width = Cm(12)
    photo.width = Cm(4.5)
    details.paragraphs[0].add_run("Hannah Keller\n")
    details.paragraphs[0].runs[0].bold = True
    set_run_font(details.paragraphs[0].runs[0], "Calibri", 13, True, "243447")
    detail_lines = [
        "Data Engineer",
        "Musterstrasse 18, 10115 Berlin",
        "hannah.keller@example.de | +49 30 555 0181",
        "github.com/synthetic-hkeller",
        "Geboren am 14.05.1994 in Bremen",
    ]
    for line in detail_lines:
        paragraph = details.add_paragraph()
        paragraph.paragraph_format.space_after = Pt(1)
        run = paragraph.add_run(line)
        set_run_font(run, "Calibri", 8.8, color="333333")
    photo.paragraphs[0].alignment = WD_ALIGN_PARAGRAPH.CENTER
    photo_run = photo.paragraphs[0].add_run("BEWERBUNGS-\nFOTO")
    set_run_font(photo_run, "Calibri", 9, True, "6E7C8C")
    set_cell_shading(photo, "E9EEF3")
    style_table(top)

    add_docx_section(doc, "BERUFSERFAHRUNG", "243447")
    work = doc.add_table(rows=4, cols=2)
    work.style = "Table Grid"
    work_data = [
        ("09/2022 - heute", "Senior Data Engineer, Nordlicht Energie, Berlin\n"
         "Aufbau einer Streaming-Plattform für 2,1 Mio. Messpunkte; Einführung "
         "automatischer Datenqualitätsprüfungen und Betriebsmetriken."),
        ("01/2020 - 08/2022", "Data Engineer, Elbe Analytics, Hamburg\n"
         "Entwicklung von Batch-Pipelines für 48 Berichte; Migration von 140 "
         "Prozessen in eine orchestrierte Cloud-Umgebung."),
        ("08/2018 - 12/2019", "Junior Data Analyst, Stadtwerke Bremen\n"
         "Analyse von Verbrauchs- und Störungsdaten; Erstellung von Dashboards "
         "für technische Teams."),
        ("03/2018 - 07/2018", "Praktikant, Institut für Energieinformatik\n"
         "Untersuchung von Anomalien in Zeitreihen mit Python und SQL."),
    ]
    for row_index, (date_text, content) in enumerate(work_data):
        date_cell, content_cell = work.rows[row_index].cells
        date_cell.width = Cm(4)
        content_cell.width = Cm(12.5)
        date_run = date_cell.paragraphs[0].add_run(date_text)
        set_run_font(date_run, "Calibri", 8.6, True, "243447")
        content_run = content_cell.paragraphs[0].add_run(content)
        set_run_font(content_run, "Calibri", 8.7, color="333333")
    style_table(work)

    add_docx_section(doc, "AUSBILDUNG", "243447")
    education = doc.add_table(rows=2, cols=2)
    education.style = "Table Grid"
    education_data = [
        ("10/2015 - 02/2018", "M.Sc. Data Engineering, Technische Universität Hamburg\n"
         "Schwerpunkt: verteilte Systeme, Datenbanken, maschinelles Lernen."),
        ("10/2012 - 09/2015", "B.Sc. Wirtschaftsinformatik, Universität Bremen\n"
         "Abschlussarbeit zur Datenqualität in Energienetzen."),
    ]
    for row_index, (date_text, content) in enumerate(education_data):
        date_cell, content_cell = education.rows[row_index].cells
        date_cell.width = Cm(4)
        content_cell.width = Cm(12.5)
        date_run = date_cell.paragraphs[0].add_run(date_text)
        set_run_font(date_run, "Calibri", 8.6, True, "243447")
        content_run = content_cell.paragraphs[0].add_run(content)
        set_run_font(content_run, "Calibri", 8.7, color="333333")
    style_table(education)

    add_docx_section(doc, "KENNTNISSE UND SPRACHEN", "243447")
    skills = doc.add_table(rows=4, cols=2)
    skills.style = "Table Grid"
    skill_data = [
        ("Technisch", "Python, SQL, Scala, Spark, Kafka, Airflow, dbt, AWS, Terraform"),
        ("Methoden", "Datenmodellierung, Datenqualität, Monitoring, CI/CD, Mentoring"),
        ("Deutsch", "Muttersprache"),
        ("Englisch", "Verhandlungssicher (C1)"),
    ]
    for row_index, (label, value) in enumerate(skill_data):
        label_cell, value_cell = skills.rows[row_index].cells
        label_cell.width = Cm(3.5)
        value_cell.width = Cm(13)
        set_cell_shading(label_cell, "E9EEF3")
        label_run = label_cell.paragraphs[0].add_run(label)
        set_run_font(label_run, "Calibri", 8.7, True, "243447")
        value_run = value_cell.paragraphs[0].add_run(value)
        set_run_font(value_run, "Calibri", 8.7, color="333333")
    style_table(skills)

    add_para(
        doc,
        "Berlin, 29.09.2026",
        size=8.5,
        color="333333",
        before=8,
        after=6,
    )
    add_para(doc, "Hannah Keller", size=10, color="333333")
    doc.save(path)


def build_bilingual_docx(path: Path) -> None:
    doc = Document()
    set_doc_defaults(doc, "Calibri", 9)
    doc.core_properties.title = "Qiao Wen - Bilingual Academic CV Fixture"
    doc.core_properties.author = "Synthetic Test Data"

    add_para(
        doc,
        "QIAO WEN  |  温乔",
        font="Calibri",
        size=20,
        bold=True,
        color="7A2E2E",
        align=WD_ALIGN_PARAGRAPH.CENTER,
        after=2,
    )
    add_para(
        doc,
        "Senior Research Scientist | Multimodal Retrieval and Evaluation",
        font="Calibri",
        size=10.5,
        color="4B4B4B",
        align=WD_ALIGN_PARAGRAPH.CENTER,
        after=2,
    )
    add_para(
        doc,
        "qiao.wen@example.edu | +86 10 5550 0188 | Beijing / Toronto | "
        "qiao-wen.example",
        size=8.5,
        color="555555",
        align=WD_ALIGN_PARAGRAPH.CENTER,
        after=5,
    )

    add_docx_section(doc, "RESEARCH INTERESTS / 研究方向", "7A2E2E")
    add_para(
        doc,
        "Multimodal retrieval, evaluation of generative systems, document understanding, "
        "cross-lingual search. 可解释评估、多模态检索、生成系统评测、跨语言搜索。",
        size=8.8,
        color="333333",
        after=3,
    )

    add_docx_section(doc, "EDUCATION / 教育经历", "7A2E2E")
    education = doc.add_table(rows=3, cols=2)
    education.style = "Table Grid"
    education_data = [
        (
            "Ph.D. Computer Science | Northbridge Institute of Technology | 2016 - 2020",
            "Advisor: Prof. Arden Lee. Dissertation: Cross-Lingual Evidence Retrieval "
            "for Scientific Question Answering.",
        ),
        (
            "M.Sc. Information Systems | Westmere University | 2013 - 2016",
            "Thesis: Evaluation Metrics for Ranked Document Summaries.",
        ),
        (
            "B.Eng. Software Engineering | Xinhai University | 2009 - 2013",
            "Graduation project: Chinese-English Patent Search.",
        ),
    ]
    for row_index, (left_text, right_text) in enumerate(education_data):
        left_cell, right_cell = education.rows[row_index].cells
        left_cell.width = Cm(6.7)
        right_cell.width = Cm(9.8)
        left_run = left_cell.paragraphs[0].add_run(left_text)
        set_run_font(left_run, "Calibri", 8.7, True, "333333")
        right_run = right_cell.paragraphs[0].add_run(right_text)
        set_run_font(right_run, "Calibri", 8.5, color="444444")
    style_table(education)

    add_docx_section(doc, "APPOINTMENTS / 任职经历", "7A2E2E")
    add_para(
        doc,
        "Senior Research Scientist, Northbridge AI Lab, 2023 - Present\n"
        "Research Scientist, Northbridge AI Lab, 2020 - 2023\n"
        "Visiting Researcher, University of Calder, 2019",
        size=8.8,
        color="333333",
        after=3,
    )

    add_docx_section(doc, "FUNDING / 科研项目", "7A2E2E")
    funding = doc.add_table(rows=3, cols=4)
    funding.style = "Table Grid"
    funding_rows = [
        ("Role", "Project", "Period", "Amount"),
        ("PI", "Reliable Multimodal Retrieval", "2025 - 2028", "CNY 680,000"),
        ("Co-I", "Cross-Lingual Evaluation Suite", "2023 - 2025", "CAD 210,000"),
    ]
    for row_index, values in enumerate(funding_rows):
        for col_index, value in enumerate(values):
            cell = funding.rows[row_index].cells[col_index]
            run = cell.paragraphs[0].add_run(value)
            set_run_font(run, "Calibri", 8.3, row_index == 0, "FFFFFF" if row_index == 0 else "333333")
            if row_index == 0:
                set_cell_shading(cell, "7A2E2E")
    style_table(funding, header_row=0, accent="7A2E2E")

    doc.add_page_break()
    add_docx_section(doc, "SELECTED PUBLICATIONS / 代表论文", "7A2E2E")
    publications = [
        "Wen, Q., Lee, A., and Gomez, P. (2025). Evidence-aware evaluation for "
        "multimodal question answering. Journal of Multimodal Research, 8(2), 101-128. "
        "Citations: 46.",
        "Wen, Q., and Zhao, M. (2024). Cross-lingual retrieval under limited parallel "
        "data. Proceedings of the Synthetic Information Retrieval Conference, 220-233. "
        "Citations: 33.",
        "Gomez, P., Wen, Q., and Lee, A. (2024). When do image captions help document "
        "retrieval? Findings of the Northbridge NLP Forum, 55-69. Citations: 19.",
        "Wen, Q. (2022). Evaluation metrics for evidence-grounded summaries. Journal of "
        "Applied Language Technology, 5(4), 410-431. Citations: 57.",
    ]
    for index, publication in enumerate(publications, 1):
        add_para(
            doc,
            f"{index}. {publication}",
            size=8.6,
            color="333333",
            after=3,
        )

    add_docx_section(doc, "ACADEMIC SERVICE / 学术服务", "7A2E2E")
    add_para(
        doc,
        "Reviewer: Synthetic Information Retrieval Conference (2023 - 2026); "
        "Journal of Applied Language Technology (2021 - Present). "
        "Program committee member: Northbridge NLP Forum (2024).",
        size=8.7,
        color="333333",
        after=3,
    )
    add_docx_section(doc, "SKILLS / 技能", "7A2E2E")
    add_para(
        doc,
        "Python, PyTorch, Hugging Face, SQL, information retrieval, multilingual NLP, "
        "evaluation design, research supervision. 英语（工作语言）、普通话（母语）。",
        size=8.7,
        color="333333",
        after=3,
    )
    doc.save(path)


def build_scanned_source_pdf(path: Path) -> None:
    """Build the clean source used to create the image-only OCR fixture."""
    c = canvas.Canvas(str(path), pagesize=A4)
    c.setTitle("Qiao Wen - Scanned CV Source")
    c.setAuthor("Synthetic Test Data")
    page_w, page_h = A4
    left = 18 * mm
    right = page_w - 18 * mm
    width = right - left

    c.setFillColor(colors.HexColor("#F6F1EA"))
    c.rect(0, 0, page_w, page_h, fill=1, stroke=0)
    c.setFillColor(colors.HexColor("#6F3B32"))
    c.rect(0, page_h - 38 * mm, page_w, 38 * mm, fill=1, stroke=0)
    c.setFillColor(colors.white)
    c.setFont("MicrosoftYaHei-Bold", 23)
    c.drawString(left, page_h - 19 * mm, "QIAO WEN  |  温乔")
    c.setFont("Arial", 10)
    c.drawString(left, page_h - 27 * mm, "Principal Research Scientist")
    c.setFont("Arial", 8.2)
    c.drawString(
        left,
        page_h - 34 * mm,
        "qiao.wen@example.edu | +86 10 5550 0188 | Beijing / Toronto",
    )

    y = page_h - 51 * mm
    c.setFillColor(colors.HexColor("#6F3B32"))
    c.setFont("MicrosoftYaHei-Bold", 11)
    c.drawString(left, y, "EDUCATION / 教育经历")
    y -= 8 * mm
    c.setFillColor(colors.HexColor("#242424"))
    c.setFont("Arial-Bold", 9)
    c.drawString(left, y, "Ph.D. Computer Science")
    c.setFont("Arial", 8.5)
    c.drawRightString(right, y, "2016 - 2021")
    y -= 5 * mm
    y = draw_wrapped(
        c,
        "Northbridge Institute of Technology | Advisor: Prof. Arden Lee\n"
        "Dissertation: Cross-Lingual Evidence Retrieval for Scientific Question Answering",
        left,
        y,
        width,
        font="Arial",
        size=8.4,
        leading=4.8 * mm,
        color=colors.HexColor("#333333"),
    )
    y -= 3 * mm

    c.setFillColor(colors.HexColor("#6F3B32"))
    c.setFont("MicrosoftYaHei-Bold", 11)
    c.drawString(left, y, "APPOINTMENTS / 任职经历")
    y -= 8 * mm
    c.setFillColor(colors.HexColor("#242424"))
    c.setFont("Arial-Bold", 9)
    c.drawString(left, y, "Principal Research Scientist")
    c.setFont("Arial", 8.5)
    c.drawRightString(right, y, "2024 - Present")
    y -= 5 * mm
    y = draw_wrapped(
        c,
        "Northbridge AI Lab | Multimodal retrieval and evaluation",
        left,
        y,
        width,
        font="Arial",
        size=8.4,
        leading=4.8 * mm,
        color=colors.HexColor("#333333"),
    )
    y -= 2 * mm
    c.setFillColor(colors.HexColor("#242424"))
    c.setFont("Arial-Bold", 9)
    c.drawString(left, y, "Research Scientist")
    c.setFont("Arial", 8.5)
    c.drawRightString(right, y, "2021 - 2024")
    y -= 5 * mm
    y = draw_wrapped(
        c,
        "Northbridge AI Lab | Cross-lingual search and document understanding",
        left,
        y,
        width,
        font="Arial",
        size=8.4,
        leading=4.8 * mm,
        color=colors.HexColor("#333333"),
    )
    y -= 5 * mm

    c.setFillColor(colors.HexColor("#6F3B32"))
    c.setFont("MicrosoftYaHei-Bold", 11)
    c.drawString(left, y, "SELECTED PUBLICATIONS / 代表论文")
    y -= 8 * mm
    publications = [
        "Wen, Q., Lee, A., and Gomez, P. (2025). Evidence-aware evaluation for "
        "multimodal question answering. Journal of Multimodal Research, 8(2), 101-128.",
        "Wen, Q., and Zhao, M. (2024). Cross-lingual retrieval under limited parallel "
        "data. Proceedings of the Synthetic Information Retrieval Conference, 220-233.",
        "Wen, Q. (2022). Evaluation metrics for evidence-grounded summaries. Journal of "
        "Applied Language Technology, 5(4), 410-431.",
    ]
    for publication in publications:
        y = draw_bullets(
            c,
            [publication],
            left,
            y,
            width,
            font="Arial",
            size=8.2,
            leading=4.7 * mm,
            color=colors.HexColor("#333333"),
        )
        y -= 1.5 * mm
    y -= 4 * mm

    c.setFillColor(colors.HexColor("#6F3B32"))
    c.setFont("MicrosoftYaHei-Bold", 11)
    c.drawString(left, y, "ACADEMIC SERVICE / 学术服务")
    y -= 8 * mm
    y = draw_bullets(
        c,
        [
            "Reviewer, Synthetic Information Retrieval Conference (2023 - 2026)",
            "Program committee member, Northbridge NLP Forum (2024)",
        ],
        left,
        y,
        width,
        font="Arial",
        size=8.2,
        leading=4.7 * mm,
        color=colors.HexColor("#333333"),
    )
    add_footer_to_static_page(c, 1)
    c.save()


def write_manifest() -> None:
    rows = [
        {
            "filename": "cv_01_academic_standard_en.pdf",
            "format": "PDF",
            "language": "English",
            "style": "Classic academic, single column, multi-page",
            "text_layer": "yes",
            "test_focus": "Long narrative, publications, grants, awards, service",
        },
        {
            "filename": "cv_02_industry_ats_en.pdf",
            "format": "PDF",
            "language": "English",
            "style": "ATS-friendly one-page industry resume",
            "text_layer": "yes",
            "test_focus": "Plain reading order, metrics, concise role history",
        },
        {
            "filename": "cv_03_modern_two_column_en.pdf",
            "format": "PDF",
            "language": "English",
            "style": "Modern two-column sidebar layout",
            "text_layer": "yes",
            "test_focus": "Reading-order recovery, sidebar sections, visual hierarchy",
        },
        {
            "filename": "cv_04_chinese_campus_cn.pdf",
            "format": "PDF",
            "language": "Chinese",
            "style": "Chinese campus recruitment resume",
            "text_layer": "yes",
            "test_focus": "Chinese text, dotted dates, campus projects, photo placeholder",
        },
        {
            "filename": "cv_05_creative_infographic_en.pdf",
            "format": "PDF",
            "language": "English",
            "style": "Creative infographic and timeline",
            "text_layer": "yes",
            "test_focus": "Graphic blocks, skill bars, nontrivial reading order",
        },
        {
            "filename": "cv_06_europass_style_en.docx",
            "format": "DOCX",
            "language": "English",
            "style": "Europass-inspired table layout",
            "text_layer": "yes",
            "test_focus": "Table extraction, language grid, structured headings",
        },
        {
            "filename": "cv_07_german_lebenslauf_de.docx",
            "format": "DOCX",
            "language": "German",
            "style": "German Lebenslauf with dated timeline",
            "text_layer": "yes",
            "test_focus": "German text, DD.MM.YYYY dates, timeline table, signature block",
        },
        {
            "filename": "cv_08_bilingual_academic_zh_en.docx",
            "format": "DOCX",
            "language": "Chinese and English",
            "style": "Bilingual academic CV",
            "text_layer": "yes",
            "test_focus": "Mixed-language sections, publications, grants, advisor data",
        },
        {
            "filename": "cv_09_scanned_image_only_zh.pdf",
            "format": "PDF",
            "language": "Chinese and English",
            "style": "Scanned image-only document",
            "text_layer": "none",
            "test_focus": "OCR boundary case and cross-document conflict with sample 08",
        },
    ]
    for row in rows:
        path = OUT_DIR / row["filename"]
        row["size_bytes"] = path.stat().st_size
        row["sha256"] = hashlib.sha256(path.read_bytes()).hexdigest()
        if path.suffix.lower() == ".pdf":
            row["pages"] = len(PdfReader(str(path)).pages)
        else:
            row["pages"] = {
                "cv_06_europass_style_en.docx": 1,
                "cv_07_german_lebenslauf_de.docx": 1,
                "cv_08_bilingual_academic_zh_en.docx": 2,
            }.get(row["filename"], "render-check")

    fieldnames = [
        "filename",
        "format",
        "language",
        "style",
        "pages",
        "text_layer",
        "size_bytes",
        "sha256",
        "test_focus",
    ]
    with (OUT_DIR / "manifest.csv").open("w", newline="", encoding="utf-8-sig") as handle:
        writer = csv.DictWriter(handle, fieldnames=fieldnames)
        writer.writeheader()
        writer.writerows(rows)

    readme = f"""# CV parser test fixtures

All people, organisations, publications, awards, and links in this folder are
fictional synthetic data. They must not be treated as real candidate claims.

Query date recorded by the fixture set: `{QUERY_DATE}`.

## Files

| File | Format | Language | Style | Text layer | Main test focus |
|---|---|---|---|---|---|
"""
    for row in rows:
        readme += (
            f"| `{row['filename']}` | {row['format']} | {row['language']} | "
            f"{row['style']} | {row['text_layer']} | {row['test_focus']} |\n"
        )
    readme += """

## Suggested upload batches

The backend contract accepts up to four PDF or DOCX files per upload.

1. `cv_01_academic_standard_en.pdf`, `cv_02_industry_ats_en.pdf`,
   `cv_03_modern_two_column_en.pdf`, `cv_04_chinese_campus_cn.pdf`
2. `cv_05_creative_infographic_en.pdf`, `cv_06_europass_style_en.docx`,
   `cv_07_german_lebenslauf_de.docx`, `cv_08_bilingual_academic_zh_en.docx`
3. `cv_09_scanned_image_only_zh.pdf` on its own

## Deliberate edge cases

- `cv_03_modern_two_column_en.pdf` and `cv_05_creative_infographic_en.pdf`
  require layout-aware reading order.
- `cv_04_chinese_campus_cn.pdf` uses Chinese names, dotted dates, and a
  photo placeholder.
- `cv_06_europass_style_en.docx` and `cv_07_german_lebenslauf_de.docx`
  contain table-heavy structures.
- `cv_07_german_lebenslauf_de.docx` uses `DD.MM.YYYY` dates.
- `cv_08_bilingual_academic_zh_en.docx` and
  `cv_09_scanned_image_only_zh.pdf` intentionally disagree on the candidate's
  PhD end date and current position. Upload them together to test
  cross-document conflict handling.
- `cv_09_scanned_image_only_zh.pdf` has no extractable text layer and should
  exercise OCR or a clear unsupported-input message.

## Integrity

`manifest.csv` records page counts, file sizes, and SHA-256 hashes.
"""
    (OUT_DIR / "README.md").write_text(readme, encoding="utf-8")


def build_scanned_pdf(source_pdf: Path, output_pdf: Path) -> None:
    pages = convert_from_path(
        str(source_pdf),
        dpi=190,
        poppler_path=str(POPPLER_DIR) if POPPLER_DIR else None,
        fmt="png",
    )
    processed: list[Image.Image] = []
    random.seed(1903)
    for page_index, page in enumerate(pages):
        page = page.convert("RGB")
        page = ImageEnhance.Contrast(page).enhance(0.88)
        page = ImageEnhance.Color(page).enhance(0.92)
        page = page.rotate(
            0.22 * (1 if page_index % 2 == 0 else -1),
            resample=Image.Resampling.BICUBIC,
            expand=True,
            fillcolor=(244, 241, 233),
        )
        canvas_image = Image.new("RGB", (page.width + 18, page.height + 18), (238, 234, 226))
        canvas_image.paste(page, (9, 9))
        noise = Image.effect_noise(canvas_image.size, 8).convert("L")
        noise = ImageOps.colorize(noise, black="#EEE8DC", white="#FFFFFF").convert("RGB")
        canvas_image = Image.blend(canvas_image, noise, 0.035)
        processed.append(canvas_image)
    first, rest = processed[0], processed[1:]
    first.save(
        output_pdf,
        "PDF",
        resolution=190,
        save_all=True,
        append_images=rest,
    )


def main() -> None:
    OUT_DIR.mkdir(parents=True, exist_ok=True)
    register_fonts()

    academic = OUT_DIR / "cv_01_academic_standard_en.pdf"
    ats = OUT_DIR / "cv_02_industry_ats_en.pdf"
    two_column = OUT_DIR / "cv_03_modern_two_column_en.pdf"
    chinese = OUT_DIR / "cv_04_chinese_campus_cn.pdf"
    infographic = OUT_DIR / "cv_05_creative_infographic_en.pdf"
    europass = OUT_DIR / "cv_06_europass_style_en.docx"
    german = OUT_DIR / "cv_07_german_lebenslauf_de.docx"
    bilingual = OUT_DIR / "cv_08_bilingual_academic_zh_en.docx"
    scanned = OUT_DIR / "cv_09_scanned_image_only_zh.pdf"

    build_academic_pdf(academic)
    build_ats_pdf(ats)
    build_two_column_pdf(two_column)
    build_chinese_campus_pdf(chinese)
    build_infographic_pdf(infographic)
    build_europass_docx(europass)
    build_german_docx(german)
    build_bilingual_docx(bilingual)
    scan_source = OUT_DIR / "_scan_source_qiao.pdf"
    build_scanned_source_pdf(scan_source)
    build_scanned_pdf(scan_source, scanned)
    scan_source.unlink()
    write_manifest()

    created = [academic, ats, two_column, chinese, infographic, europass, german, bilingual, scanned]
    for path in created:
        print(f"{path.name}\t{path.stat().st_size} bytes")


if __name__ == "__main__":
    main()
